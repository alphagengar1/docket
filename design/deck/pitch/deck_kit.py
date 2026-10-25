#!/usr/bin/env python3
"""Shared layout kit for the editable decks (build_native.py, build_benchmarks.py): native text boxes,
tables and charts in python-pptx, Arial + Courier New, Instrument Tag colours. Import with
`from deck_kit import *`; each build script sets deck_kit.NOTES or passes notes= to new_slide.
"""
import json
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from lxml import etree

HERE = Path(__file__).resolve().parent
NOTES = []  # set by the build script

INK, MUTE, HAIR, PANEL, PAPER = "141514", "565952", "C9CBC4", "F5F5F1", "FFFFFF"
SIGNAL, GO = "FF5A1F", "1C46C9"
SANS, MONO = "Arial", "Courier New"
W, H = 13.333, 7.5
L, R = 0.67, 0.67           # side margins (inches)
CW = W - L - R              # content width

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(W), Inches(H)
BLANK = prs.slide_layouts[6]

def rgb(h): return RGBColor.from_string(h)

# ---------- primitives ----------
def box(slide, x, y, w, h, fill=None, line=None, line_w=0.75, dashed=False, shape=MSO_SHAPE.RECTANGLE):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    s.shadow.inherit = False
    if fill: s.fill.solid(); s.fill.fore_color.rgb = rgb(fill)
    else: s.fill.background()
    if line:
        s.line.color.rgb = rgb(line); s.line.width = Pt(line_w)
        if dashed:
            s.line.dash_style = 4  # MSO_LINE.DASH
    else: s.line.fill.background()
    return s

def text(slide, x, y, w, h, runs, size=14, color=INK, font=SANS, bold=False, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, line_spacing=1.15, margin=0, name=None):
    """runs: str, or list of paragraphs; a paragraph is a str or a list of (text, {opts}) runs."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    if name: tb.name = name
    tf = tb.text_frame; tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = Inches(margin)
    tf.vertical_anchor = anchor
    paras = runs if isinstance(runs, list) else [runs]
    for i, p in enumerate(paras):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = align; para.line_spacing = line_spacing
        for r in (p if isinstance(p, list) else [(p, {})]):
            t, o = (r, {}) if isinstance(r, str) else r
            run = para.add_run(); run.text = t
            f = run.font; f.size = Pt(o.get("size", size)); f.bold = o.get("bold", bold)
            f.name = o.get("font", font); f.color.rgb = rgb(o.get("color", color))
    return tb

def label(slide, x, y, w, t, size=9, color=MUTE, h=0.25):
    return text(slide, x, y, w, h, t.upper(), size=size, color=color, font=MONO)

def chip(slide, x, y, t, fill=None, line=INK, color=INK, size=9, w=None, h=0.3, dashed=False, font=MONO):
    w = w or (len(t) * size * 0.6 / 72 + 0.26)   # Courier New is 0.6 em per character
    s = box(slide, x, y, w, h, fill=fill, line=line, dashed=dashed)
    tf = s.text_frame; tf.margin_left = tf.margin_right = Inches(0.08); tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = t.upper() if font == MONO else t
    r.font.size = Pt(size); r.font.name = font; r.font.color.rgb = rgb(color); r.font.bold = font != MONO
    return s, w

def hline(slide, x, y, w, color=HAIR, weight=0.75):
    c = slide.shapes.add_connector(1, Inches(x), Inches(y), Inches(x + w), Inches(y))
    c.line.color.rgb = rgb(color); c.line.width = Pt(weight)
    return c

def _border(cell, side, color, w_pt):
    tcPr = cell._tc.get_or_add_tcPr()
    for old in tcPr.findall(qn(f"a:{side}")): tcPr.remove(old)
    ln = etree.SubElement(tcPr, qn(f"a:{side}"), w=str(int(w_pt * 12700)), cap="flat", cmpd="sng", algn="ctr")
    sf = etree.SubElement(ln, qn("a:solidFill")); etree.SubElement(sf, qn("a:srgbClr"), val=color)
    etree.SubElement(ln, qn("a:prstDash"), val="solid")

def _noborder(cell, side):
    tcPr = cell._tc.get_or_add_tcPr()
    for old in tcPr.findall(qn(f"a:{side}")): tcPr.remove(old)
    ln = etree.SubElement(tcPr, qn(f"a:{side}"), w="0"); etree.SubElement(ln, qn("a:noFill"))

def table(slide, x, y, w, cols, rows, size=12, head_size=8.5, row_h=0.42, head_h=0.34):
    """cols: list of (width_in, kind) where kind in {'sans','mono'}; rows[0] is the header.
    Cell values: str or list of (text, opts) runs."""
    shape = slide.shapes.add_table(len(rows), len(cols), Inches(x), Inches(y), Inches(w), Inches(head_h + row_h * (len(rows) - 1)))
    tbl = shape.table
    tblPr = tbl._tbl.tblPr
    tblPr.set("firstRow", "0"); tblPr.set("bandRow", "0")
    style = tblPr.find(qn("a:tableStyleId"))
    if style is None: style = etree.SubElement(tblPr, qn("a:tableStyleId"))
    style.text = "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"  # No Style, No Grid
    for ci, (cw, _) in enumerate(cols): tbl.columns[ci].width = Inches(cw)
    for ri, row in enumerate(rows):
        tbl.rows[ri].height = Inches(head_h if ri == 0 else row_h)
        for ci, val in enumerate(row):
            cell = tbl.cell(ri, ci); cell.fill.background()
            cell.margin_left = Inches(0); cell.margin_right = Inches(0.1)
            cell.margin_top = Inches(0.05 if ri else 0.0); cell.margin_bottom = Inches(0.05)
            cell.vertical_anchor = MSO_ANCHOR.TOP if ri else MSO_ANCHOR.BOTTOM
            tf = cell.text_frame; tf.word_wrap = True
            p = tf.paragraphs[0]
            # A cell is a str, a list of (text, opts) runs, or a list of paragraphs (each a list of runs).
            if isinstance(val, list) and val and isinstance(val[0], list):
                paras = val
            else:
                paras = [val if isinstance(val, list) else [(val, {})]]
            for pi, runs in enumerate(paras):
                if pi:
                    p = tf.add_paragraph(); p.space_before = Pt(2)
                for t, o in runs:
                    r = p.add_run(); r.text = (t.upper() if ri == 0 else t)
                    mono = ri == 0 or cols[ci][1] == "mono"
                    r.font.name = o.get("font", MONO if mono else SANS)
                    r.font.size = Pt(o.get("size", head_size if ri == 0 else (size - 1 if mono else size)))
                    r.font.bold = o.get("bold", False)
                    r.font.color.rgb = rgb(o.get("color", MUTE if ri == 0 else INK))
            for side in ("lnL", "lnR", "lnT"): _noborder(cell, side)
            if ri == 0: _border(cell, "lnB", INK, 1.25)
            elif ri < len(rows) - 1: _border(cell, "lnB", HAIR, 0.75)
            else: _noborder(cell, "lnB")
    return shape

# ---------- page furniture ----------
STEPS = ["1 Data", "2 Hard stops", "3 Underwriters ×5", "4 Verify", "5 Notice", "6 Human + platform"]

def wordmark(slide, x, cy, pt):
    """Orange square + DOCKET, vertically centred on cy; the square matches the letters' cap height."""
    side = 0.75 * pt / 72                     # Arial Bold cap height is ~0.72 em; a hair taller reads as equal
    gap = side * 0.45
    box(slide, x, cy - side / 2, side, side, fill=SIGNAL)
    h = pt * 1.4 / 72
    tb = text(slide, x + side + gap, cy - h / 2, pt * 0.11 * 6, h, "DOCKET", size=pt, bold=True, font=SANS,
              anchor=MSO_ANCHOR.MIDDLE, line_spacing=1.0, margin=0)
    tb.text_frame.word_wrap = False
    return tb

def frame(slide, title, active=0, number=None, source=None, tracker=True):
    wordmark(slide, L, 0.41, 14)
    # tracker, right-aligned
    widths = []
    if not tracker: STEPS_ = []
    else: STEPS_ = STEPS
    widths = [len(s) * 8 * 0.6 / 72 + 0.3 for s in STEPS_]
    x = W - R - sum(widths) - 0.05 * (len(STEPS) - 1)
    for i, (s, cw) in enumerate(zip(STEPS_, widths)):
        on = (i + 1) == active
        chip(slide, x, 0.28, s, fill=INK if on else None, line=INK if on else HAIR, color=PAPER if on else MUTE, size=8, w=cw, h=0.26)
        x += cw + 0.05
    # action title
    text(slide, L, 0.72, CW, 1.0, title, size=24, bold=True, anchor=MSO_ANCHOR.TOP, line_spacing=1.1, name="Title")
    # footer
    if source: text(slide, L, H - 0.62, CW - 1, 0.45, source, size=8, color=MUTE, font=MONO, anchor=MSO_ANCHOR.BOTTOM, name="Source")
    # page number box is always present (named PageNum) so a combined deck can renumber after reordering
    text(slide, W - R - 0.6, H - 0.5, 0.6, 0.3, str(number) if number else "", size=8, color=MUTE, font=MONO, align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.BOTTOM, name="PageNum")

def takeaway(slide, x, y, w, runs, size=15, h=None):
    """Key message in a tinted panel."""
    lines = max(1, int(len("".join(t for t, _ in runs)) * size * 0.0075 / w) + 1)
    h = h or (0.2 + lines * size * 1.3 / 72 + 0.15)
    box(slide, x, y, w, h, fill=PANEL).name = "Takeaway"
    text(slide, x + 0.22, y + 0.12, w - 0.44, h - 0.24, [runs], size=size, anchor=MSO_ANCHOR.MIDDLE, line_spacing=1.2, name="TakeawayText")
    return h

def chip_row(slide, x, y, max_x, items, size=8, h=0.28, gap=0.1):
    """Chips left to right, wrapping to a new row before max_x. items: (text, fill, line, color, dashed)."""
    xx, yy = x, y
    for t, fill, line, color, dashed in items:
        cw = len(t) * size * 0.6 / 72 + 0.26
        if xx + cw > max_x and xx > x:
            xx, yy = x, yy + h + 0.1
        chip(slide, xx, yy, t, fill=fill, line=line, color=color, size=size, w=cw, h=h, dashed=dashed)
        xx += cw + gap
    return yy + h

def card(slide, x, y, w, h, lab, runs, size=12, dashed=False, fill=PANEL):
    box(slide, x, y, w, h, fill=None if dashed else fill, line=HAIR, dashed=dashed)
    label(slide, x + 0.2, y + 0.14, w - 0.4, lab)
    text(slide, x + 0.2, y + 0.42, w - 0.4, h - 0.5, runs, size=size, line_spacing=1.2)

def bullets(slide, x, y, w, h, items, size=12):
    tb = text(slide, x, y, w, h, [[(t, {})] if isinstance(t, str) else t for t in items], size=size, line_spacing=1.15)
    for p in tb.text_frame.paragraphs:
        pPr = p._p.get_or_add_pPr(); pPr.set("marL", str(int(Inches(0.22)))); pPr.set("indent", str(-int(Inches(0.22))))
        bu = etree.SubElement(pPr, qn("a:buChar")); bu.set("char", "•")
        p.space_after = Pt(4)
    return tb

def new_slide(i, hidden=False, notes=None):
    s = prs.slides.add_slide(BLANK)
    s.notes_slide.notes_text_frame.text = notes if notes is not None else NOTES[i - 1]
    if hidden: s._element.set("show", "0")
    return s

B = {"bold": True}
M = {"font": MONO}

