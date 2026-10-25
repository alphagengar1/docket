#!/usr/bin/env python3
"""Build the v7 decks: one source, two talk orders.

  python3 design/deck/pitch/build_v7.py general    -> docket-deck-v7-general.pptx
  python3 design/deck/pitch/build_v7.py technical  -> docket-deck-v7-technical.pptx
  python3 design/deck/pitch/build_v7.py            -> both (each in its own process)

Both files hold every slide in the same order; each version hides the slides it skips, so either file can
be switched to the other talk by unhiding. Adds a market slide (sourced dollar figure in the first 10
seconds) and a business-model slide to v6. Market sources are in the speaker notes; pricing is a proposal.

Shared order:
  cover · market · problem · vs_humans · speed · stack · l1–l6 · (creao, always hidden) · eval · moat ·
  business · close · backup · scorecard · method
  general   hides l1–l6 (the demo runs on the six-layer slide)
  technical hides vs_humans and speed
"""
import re, runpy, subprocess, sys
VARIANT = sys.argv[1] if len(sys.argv) > 1 else None
if VARIANT is None:
    for v in ("general", "technical"):
        subprocess.run([sys.executable, __file__, v], check=True)
    sys.exit(0)
assert VARIANT in ("general", "technical"), VARIANT

import deck_kit
from deck_kit import *
from pptx.util import Emu

OUT = HERE / f"docket-deck-v7-{VARIANT}.pptx"

deck_kit.prs.save = lambda *a, **k: None
runpy.run_path(str(HERE / "build_native.py"), run_name="v7_main")       # slides 0-11
runpy.run_path(str(HERE / "build_benchmarks.py"), run_name="v7_bench")  # slides 12-16

# 17 Backup divider
s = new_slide(0, notes="BACKUP. Everything after this slide is for questions, not for the talk.")
box(s, L, 3.1, 0.5, 0.5, fill=SIGNAL)
text(s, L + 0.75, 2.95, 9, 0.8, "Backup", size=40, bold=True)
text(s, L + 0.75, 3.75, 9, 0.5, "Details for questions: the agent scorecard and how every number was measured.", size=14, color=MUTE)

# 18 Market
s = new_slide(0, notes=(
    "(A) $281 billion in personal loans outstanding, a record, and about one in six people who apply for credit get turned down. "
    "Every one of those gets a letter, and the letter tells them almost nothing. That's what's broken.\n"
    "IF ASKED: $281B and 26.9M borrowers: TransUnion Q2 2026 Credit Industry Insights Report (record, up 9.6% year over year). "
    "1 in 6: New York Fed SCE Credit Access Survey, June 2026, 16.1% of credit applicants rejected in the past 12 months (any kind of credit). "
    "4,214 federally insured credit unions: NCUA Quarterly Credit Union Data Summary, Q2 2026."))
frame(s, "$281 billion in U.S. personal loans, and 1 in 6 people who apply for credit get turned down", tracker=False)
tiles = [("$281B", "Personal loans owed", "A record, up 9.6% in a year"),
         ("26.9M", "Americans carrying one", "Unsecured personal loans"),
         ("1 in 6", "Applicants turned down", "16.1% of credit applicants, past year"),
         ("4,214", "Credit unions", "Our first customers")]
tw = (CW - 3 * 0.25) / 4
for k, (n, t, d) in enumerate(tiles):
    bx = L + k * (tw + 0.25)
    box(s, bx, 1.95, tw, 2.1, fill=PANEL if k == 0 else None, line=HAIR)
    text(s, bx + 0.22, 2.08, tw - 0.4, 0.85, n, size=40, bold=True, font=MONO, color=GO if k in (0, 2) else INK)
    text(s, bx + 0.22, 3.0, tw - 0.4, 0.35, t, size=14, bold=True)
    text(s, bx + 0.22, 3.34, tw - 0.4, 0.4, d, size=10, color=MUTE)
text(s, L, 4.4, CW, 0.9, [[("Every one of those turn-downs needs a letter, and most letters say nothing.", B)],
                          [("Federal law requires the real reasons. Too many letters say \"internal standards\" instead. That's what's broken.", {"size": 11, "color": MUTE})]],
     size=16, line_spacing=1.25)
takeaway(s, L, 5.75, CW, [("A record lending market, and the no is the part nobody has fixed.", B)], size=14, h=0.55)

# 19 Business model
s = new_slide(0, notes="BUSINESS")   # notes set per version below
frame(s, "The business: credit unions pay per decided application, and every new lender is just a new policy file", tracker=False)
cw3 = (CW - 2 * 0.3) / 3
cols = [("Who pays", [("Credit unions and community banks", "4,214 federally insured credit unions in the U.S."),
                      ("The loan officer is the user", "One person runs the desk; Docket does the legwork")]),
        ("What they pay (proposed)", [("$5 per decided application", "Tested reasons and the notice included"),
                                      ("500 applications a month = $2,500", "Bureau pull passed through at cost")]),
        ("Why it scales", [("Model cost about 3¢ per application", "Under 1% of the price"),
                           ("A new lender is a new policy file", "Same agents, same tests, no rebuild")])]
for k, (lab, items) in enumerate(cols):
    bx = L + k * (cw3 + 0.3)
    box(s, bx, 1.95, cw3, 2.75, fill=PANEL if k == 1 else None, line=HAIR)
    label(s, bx + 0.22, 2.08, cw3 - 0.4, lab)
    yy = 2.42
    for h_, d in items:
        text(s, bx + 0.22, yy, cw3 - 0.44, 0.32, h_, size=14, bold=True, color=GO if k == 1 and yy < 2.5 else INK)
        text(s, bx + 0.22, yy + 0.34, cw3 - 0.44, 0.4, d, size=10, color=MUTE)
        yy += 1.0
text(s, L, 4.95, CW, 0.6, [[("Next: a lender network.", B), (" Outside lenders plug in their own policy file, and we earn a referral fee on funded loans.", {"color": MUTE, "size": 12})]],
     size=14)
takeaway(s, L, 5.75, CW, [("Cents to run, dollars to sell, and every new customer is just a policy file.", B)], size=14, h=0.55)

M = {"cover": 0, "problem": 1, "stack": 2, "l1": 3, "l2": 4, "l3": 5, "l4": 6, "l5": 7, "l6": 8, "creao": 9, "eval": 10, "close": 11,
     "vs_humans": 12, "scorecard": 13, "speed": 14, "moat": 15, "method": 16, "backup": 17, "market": 18, "business": 19}
ORDER = ["cover", "market", "problem", "vs_humans", "speed", "stack", "l1", "l2", "l3", "l4", "l5", "l6", "creao",
         "eval", "moat", "business", "close", "backup", "scorecard", "method"]
HIDE = {"general": {"l1", "l2", "l3", "l4", "l5", "l6", "creao"}, "technical": {"vs_humans", "speed", "creao"}}[VARIANT]
TIME = {
    "general": {"cover": "0:00", "market": "0:05", "problem": "0:20", "vs_humans": "0:30", "speed": "0:45", "stack": "0:55",
                "eval": "1:45", "moat": "2:00", "business": "2:15", "close": "2:45"},
    "technical": {"cover": "0:00", "market": "0:05", "problem": "0:15", "stack": "0:25", "l1": "0:35", "l2": "0:45", "l3": "1:05",
                  "l4": "1:25", "l5": "1:40", "l6": "1:50", "eval": "2:05", "moat": "2:20", "business": "2:35", "close": "2:50"},
}[VARIANT]
SPOKEN = {
    "vs_humans": "(A) Why agents and not more loan officers? Read the bold lines: days become a minute, the same file gets the same answer every time, time of day doesn't matter, every reason is tested, and it costs cents. A person still signs.",
    "speed": "(A) Speed is the biggest reason. About a minute instead of 1 to 7 business days, and a fraud file stops in 5 seconds. The officer reviews instead of assembling. Now let's show you.",
    "moat": "(B) Our data comes from CRS: all three bureaus in days, not months. Anyone can call CRS, so that's our head start. The moat is the tested decisions we build on top, and it grows with every file a credit union runs through Docket.",
    "business": {
        "general": "(B) Who pays: credit unions and community banks. What they pay: we propose $5 per decided application, with the tested reasons and the notice included; a credit union deciding 500 a month pays $2,500. Model cost is about 3 cents, under 1% of the price. And every new lender is just a new policy file, same agents. Next: outside lenders plug in, and we earn a referral fee on funded loans.",
        "technical": "(B) One line: credit unions pay per decided application, model cost is under 1% of the price, and a new lender is a new policy file. Move on.",
    }[VARIANT],
}
STACK_EXTRA = {
    "general": " GENERAL VERSION: the layer slides are hidden. Stay on this slide for the whole demo (card code, credit pull, hard-stop card and override, offers, approve), pointing at each row as it happens, then go to the eval.",
    "technical": " TECHNICAL VERSION: click through the six layer slides as the demo reaches each step.",
}[VARIANT]

slides = list(prs.slides)
assert len(slides) == 20, len(slides)
lst = prs.slides._sldIdLst
ids = list(lst)
for el in ids: lst.remove(el)
for key in ORDER: lst.append(ids[M[key]])

# Hide per version; number only the slides that are shown.
n = 0
for key in ORDER:
    sl = slides[M[key]]
    if key in HIDE: sl._element.set("show", "0")
    elif sl._element.get("show") == "0": del sl._element.attrib["show"]
    if key not in HIDE: n += 1
    for sh in sl.shapes:
        if sh.name == "PageNum" and sh.has_text_frame:
            lab = "" if key in ("cover", "backup") or key in HIDE else str(n)
            runs = sh.text_frame.paragraphs[0].runs
            if runs: runs[0].text = lab
            else: sh.text_frame.paragraphs[0].text = lab

for key in ORDER:
    tf = slides[M[key]].notes_slide.notes_text_frame
    body = tf.text
    if key in SPOKEN:
        rest = body.split("IF ASKED:", 1)
        body = SPOKEN[key] + ("\nIF ASKED:" + rest[1] if len(rest) > 1 else "")
    if key == "stack":
        body = body.split("\nIF ASKED:", 1)
        body = body[0] + STACK_EXTRA + ("\nIF ASKED:" + body[1] if len(body) > 1 else "")
    if key in TIME:
        body = f"{TIME[key]} " + re.sub(r"^\d:\d\d\s*", "", body)
    elif key in HIDE and key.startswith("l"):
        body = "HIDDEN IN THE GENERAL VERSION (the demo runs on the six-layer slide). " + re.sub(r"^\d:\d\d\s*", "", body)
    elif key in HIDE and key in ("vs_humans", "speed"):
        body = "HIDDEN IN THE TECHNICAL VERSION. " + re.sub(r"^\d:\d\d\s*", "", body)
    elif key in ("scorecard", "method"):
        body = "BACKUP. " + re.sub(r"^BACKUP[^.]*\.\s*", "", body)
    elif key == "creao":
        body = re.sub(r"ALTERNATE for slide \d+ \(hidden\): unhide this one and hide slide \d+",
                      "ALTERNATE for the Layer 6 slide just before it (hidden): unhide this one and hide that one", body)
    tf.text = body

from deck_fit import fit_slide, overflow_report
for key in ORDER:
    sl = slides[M[key]]
    fit_slide(sl)
    if key == "close":
        pic = next(sh for sh in sl.shapes if sh.shape_type == 13)
        panel = next(sh for sh in sl.shapes if sh.shape_type == 1 and sh.left <= pic.left and sh.left + sh.width >= pic.left + pic.width and sh.top <= pic.top + pic.height and sh.top + sh.height >= pic.top)
        cap = next(sh for sh in sl.shapes if sh.has_text_frame and "BE THE BORROWER" in sh.text_frame.text)
        side = min(Emu(panel.height).inches - 0.75, 4.2)
        pw = side + 0.4; right = Emu(panel.left + panel.width).inches
        panel.left, panel.width = Inches(right - pw), Inches(pw)
        pic.left, pic.top, pic.width, pic.height = Inches(right - pw + 0.2), Emu(panel.top + Inches(0.2)), Inches(side), Inches(side)
        cap.left, cap.width, cap.top = panel.left, panel.width, Emu(pic.top + pic.height + Inches(0.12))
for issue in overflow_report(prs):
    if issue[2] != "content runs past the bottom" or issue[3] > 7.4: print("FIT?", issue)

type(prs).save(prs, OUT)
shown = [k for k in ORDER if k not in HIDE]
print(f"wrote {OUT.name}: {len(ORDER)} slides, {len(shown)} shown -> {' · '.join(shown)}")
